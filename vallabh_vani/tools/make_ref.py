#!/usr/bin/env python3
import subprocess
subprocess.run("pandoc --print-default-data-file reference.docx > ref_base.docx", shell=True, check=True)
from docx import Document
from docx.shared import Pt, Mm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
doc = Document("ref_base.docx")
MAROON="7B2D26"; TEAL="2C4A52"; GREY="555049"
def strip_theme(rf):
    # pandoc's reference.docx puts theme-font attributes (w:asciiTheme etc.) on Title/Heading styles;
    # Word and LibreOffice let those override w:ascii, so the headings fell back to a sans font.
    for k in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rf.get(qn(k)) is not None: del rf.attrib[qn(k)]
def rfonts(style, a, cs):
    rpr = style.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    strip_theme(rf)
    rf.set(qn('w:ascii'), a); rf.set(qn('w:hAnsi'), a); rf.set(qn('w:cs'), cs); rf.set(qn('w:eastAsia'), cs)
def setstyle(name, size, color=None, bold=None, italic=None, a="Noto Serif", cs="Noto Serif Devanagari"):
    st = next(s for s in doc.styles if s.name == name)
    st.font.name = a; st.font.size = Pt(size)
    if color: st.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: st.font.bold = bold
    if italic is not None: st.font.italic = italic
    rfonts(st, a, cs); return st
n = setstyle("Normal", 11); n.paragraph_format.line_spacing = 1.12; n.paragraph_format.space_after = Pt(8)
setstyle("Heading 1", 17, MAROON, bold=True); setstyle("Heading 2", 13.5, MAROON, bold=True)
setstyle("Heading 3", 12, MAROON, bold=True); setstyle("Title", 24, MAROON, bold=True)
setstyle("Subtitle", 13, TEAL, italic=True)
bt = setstyle("Block Text", 10.5); bt.paragraph_format.left_indent = Inches(0.55)
bt.paragraph_format.right_indent = Inches(0.3); bt.paragraph_format.space_after = Pt(4)
cap = setstyle("Image Caption", 9, GREY, italic=True); cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
sec = doc.sections[0]; sec.page_width = Mm(210); sec.page_height = Mm(297)
sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; run = p.add_run()
for t in ('begin', None, 'end'):
    if t: e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), t)
    else: e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = ' PAGE '
    run._r.append(e)
run.font.size = Pt(9)
for rf in doc.styles.element.iter(qn('w:rFonts')):
    if rf.get(qn('w:ascii')) is None:   # docDefaults and unset styles: name the house fonts explicitly
        strip_theme(rf); rf.set(qn('w:ascii'), "Noto Serif"); rf.set(qn('w:hAnsi'), "Noto Serif")
        rf.set(qn('w:cs'), "Noto Serif Devanagari"); rf.set(qn('w:eastAsia'), "Noto Serif Devanagari")
import sys
doc.save(sys.argv[1] if len(sys.argv) > 1 else "build/ref.docx")
