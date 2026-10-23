#!/usr/bin/env python3
"""Write the Word reference documents for the two booklet types.

    make_ref_booklet.py hindi   OUT.docx   # pravachan booklet, after प्रतिबिम्बवाद (2026, rev. 2)
    make_ref_booklet.py english OUT.docx   # translation book, after नलकूबरमणिग्रीवकृतस्तुति (rev. 1)

Every number below was measured in the exemplar PDFs (page size, margins, fonts, sizes, spacing);
see vallabh_vani/booklets/SPEC.md. pandoc uses the styles, the page setup and the footer of this file.
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

KIND, OUT = sys.argv[1], Path(sys.argv[2])
with tempfile.TemporaryDirectory() as tmp:
    base = Path(tmp) / "ref_base.docx"
    base.write_bytes(
        subprocess.run(
            ["pandoc", "--print-default-data-file", "reference.docx"],
            capture_output=True,
            check=True,
        ).stdout
    )
    doc = Document(str(base))

FONT = {"hindi": "Noto Sans Medium", "english": "Sanskrit Text"}[KIND]
LANG = {"hindi": ("en-IN", "hi-IN"), "english": ("en-GB", "hi-IN")}[KIND]


def set_fonts(rpr, name):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(k)) is not None:
            del rf.attrib[qn(k)]
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(k), name)


def child(parent, tag, **attrs):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    for k, v in attrs.items():
        el.set(qn(k), str(v))
    return el


def style(name, base="Normal", kind="paragraph"):
    """Find a style by its id (pandoc's built-in names differ in case), else add it."""
    sid = name.replace(" ", "")
    for st in doc.styles:
        if st.style_id == sid:
            return st
    st = doc.styles.add_style(name, 1 if kind == "paragraph" else 2)
    st.base_style = style(base) if base else None
    st.quick_style = True
    return st


def para(
    name,
    *,
    size=None,
    size_cs=None,
    bold=None,
    underline=None,
    italic=None,
    align=None,
    first=None,
    left=None,
    hanging=None,
    before=None,
    after=None,
    line=None,
    rule="auto",
    page_break=None,
    keep_next=None,
    tabs=(),
    base="Normal",
    font=None,
):
    st = style(name, base)
    ppr = st.element.get_or_add_pPr()
    rpr = st.element.get_or_add_rPr()
    if font:
        set_fonts(rpr, font)
    if size is not None:
        child(rpr, "w:sz", **{"w:val": int(size * 2)})
        child(rpr, "w:szCs", **{"w:val": int((size_cs or size) * 2)})
    if bold is not None:
        for t in ("w:b", "w:bCs"):
            child(rpr, t, **{"w:val": "1" if bold else "0"})
    if italic is not None:
        for t in ("w:i", "w:iCs"):
            child(rpr, t, **{"w:val": "1" if italic else "0"})
    if underline is not None:
        child(rpr, "w:u", **{"w:val": "single" if underline else "none"})
    if align:
        st.paragraph_format.alignment = {
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }[align]
    pf = st.paragraph_format
    if first is not None:
        pf.first_line_indent = Pt(first)
    if left is not None:
        pf.left_indent = Pt(left)
    if hanging is not None:
        pf.first_line_indent = Pt(-hanging)
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        sp = child(ppr, "w:spacing")
        sp.set(qn("w:line"), str(int(line * 240) if rule == "auto" else int(line * 20)))
        sp.set(qn("w:lineRule"), rule)
    if page_break is not None:
        child(ppr, "w:pageBreakBefore", **{"w:val": "1" if page_break else "0"})
    if keep_next is not None:
        child(ppr, "w:keepNext", **{"w:val": "1" if keep_next else "0"})
    if tabs:
        old = ppr.find(qn("w:tabs"))
        if old is not None:
            ppr.remove(old)
        el = OxmlElement("w:tabs")
        for pos, kind in tabs:
            t = OxmlElement("w:tab")
            t.set(qn("w:val"), kind)
            t.set(qn("w:pos"), str(int(pos * 20)))
            el.append(t)
        ppr.append(el)
    return st


# ---- document defaults: one font for both scripts, language pair, no theme fonts
rpr_default = (
    doc.styles.element.find(qn("w:docDefaults"))
    .find(qn("w:rPrDefault"))
    .find(qn("w:rPr"))
)
set_fonts(rpr_default, FONT)
lang = child(rpr_default, "w:lang")
lang.set(qn("w:val"), LANG[0])
lang.set(qn("w:bidi"), LANG[1])
for rf in doc.styles.element.iter(qn("w:rFonts")):
    set_fonts(rf.getparent(), FONT)
for st in doc.styles:  # pandoc's headings are coloured and in a theme font: neutralise
    rpr = st.element.find(qn("w:rPr"))
    if rpr is not None:
        c = rpr.find(qn("w:color"))
        if c is not None:
            rpr.remove(c)

sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(5.5), Inches(8.5)

if KIND == "hindi":
    # body: Noto Sans Medium, Devanagari 14 pt / Latin 13 pt, Word "1.15 lines", 12 pt after, justified, 0.5 in indent
    # Word lays this out at 22.1 pt per line ("1.15 lines" of Noto Sans Medium 14 pt); an exact 22.1 pt
    # gives the same pitch in Word and in LibreOffice, so pages break where the exemplar's do
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        para(
            name,
            size=13,
            size_cs=14,
            align="both",
            first=36,
            before=0,
            after=12,
            line=22.1,
            rule="exact",
        )
    # The exemplar's bottom margin is 1 in, but Word lets the last line of a page reach the margin with its
    # glyphs (19.1 pt) rather than its full 22.1 pt line; with exact line spacing the full line must fit, so the
    # margin is 3 pt smaller (69 pt): pages then hold exactly the lines Word gave them (21 with two paragraph
    # breaks, 22 with none).
    sec.top_margin, sec.bottom_margin = Inches(0.75), Pt(69)
    sec.left_margin, sec.right_margin = Inches(0.75), Inches(0.6)
    sec.footer_distance, sec.header_distance = Inches(0.5), Inches(0.5)
    para(
        "Title",
        size=36,
        bold=True,
        align="center",
        first=0,
        before=48,
        after=0,
        line=1.15,
    )
    para(
        "Subtitle",
        size=20,
        bold=True,
        italic=False,
        align="center",
        first=0,
        before=11,
        after=0,
        line=1.15,
    )
    para("Occasion", size=16, align="center", first=0, before=86, after=0, line=1.15)
    para("Speaker", size=16, align="center", first=0, before=124, after=12, line=1.15)
    para(
        "Heading 1",
        size=16,
        bold=True,
        underline=True,
        align="center",
        first=0,
        before=112.7,
        after=77.6,
        line=1.15,
        page_break=True,
        keep_next=True,
    )
    para(
        "Heading 2",
        size=18,
        bold=True,
        align="center",
        first=0,
        before=34,
        after=43.4,
        line=1.15,
        keep_next=True,
    )
    para(
        "Contents Heading",
        size=13,
        bold=True,
        align="left",
        first=0,
        before=32,
        after=12,
        line=1.15,
    )
    para(
        "Contents Item",
        size=13,
        align="both",
        left=36,
        hanging=18,
        after=12,
        line=1.15,
        tabs=[(36, "left")],
    )
    para(
        "Verse",
        size=13,
        size_cs=14,
        align="both",
        first=0,
        after=12,
        line=22.1,
        rule="exact",
    )
    para(
        "Verse Prose",
        size=13,
        size_cs=14,
        align="both",
        first=0,
        after=12,
        line=22.1,
        rule="exact",
    )
    para(
        "Colophon",
        size=13,
        size_cs=14,
        align="center",
        first=0,
        after=12,
        line=22.1,
        rule="exact",
    )
    para(
        "Pada Heading",
        size=13,
        size_cs=14,
        bold=True,
        first=36,
        after=12,
        line=22.1,
        rule="exact",
    )
    para(
        "Pada",
        size=13,
        size_cs=14,
        align="left",
        first=36,
        after=12,
        line=22.1,
        rule="exact",
    )
    para("Figure", align="center", first=0, after=12, line=1.0)
    para("Chart Text", size=11, align="left", first=0, after=0, line=1.15)
    para(
        "Index Heading",
        size=14,
        bold=True,
        underline=True,
        align="center",
        first=0,
        before=136,
        after=12,
        line=1.15,
        keep_next=True,
    )
    para(
        "Index Item",
        size=11,
        align="left",
        left=54,
        hanging=36,
        after=3,
        line=1.15,
        tabs=[(54, "left"), (290, "right")],
    )
    q = style("Quotation", base=None, kind="character")
    qr = q.element.get_or_add_rPr()
    for t in ("w:b", "w:bCs"):
        child(qr, t)
    child(qr, "w:sz", **{"w:val": 26})
    child(qr, "w:szCs", **{"w:val": 26})
    folio_font, folio_size, folio_bold = "Calibri", 12, False
else:
    # body: Sanskrit Text 12 pt on a 20 pt line, 4.3 pt between paragraphs, justified, 15 mm indent
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        para(
            name,
            size=12,
            align="both",
            first=42.5,
            before=0,
            after=4.3,
            line=20,
            rule="exact",
        )
    sec.top_margin, sec.bottom_margin = Pt(40), Pt(47)
    sec.left_margin, sec.right_margin = Pt(36), Pt(32)
    sec.footer_distance, sec.header_distance = Pt(16), Pt(20)
    para(
        "Title",
        size=21,
        align="center",
        first=0,
        before=120,
        after=4,
        line=24.3,
        rule="exact",
    )
    para(
        "Subtitle",
        size=16,
        align="center",
        first=0,
        before=0,
        after=4,
        line=24.3,
        rule="exact",
    )
    para(
        "Book Title",
        size=21,
        align="center",
        first=0,
        before=120,
        after=4,
        line=26,
        rule="exact",
    )
    para(
        "Book Title IAST",
        size=16,
        align="center",
        first=0,
        after=4,
        line=24.3,
        rule="exact",
    )
    para(
        "Book Author",
        size=18,
        align="center",
        first=0,
        before=150,
        after=4,
        line=24.3,
        rule="exact",
    )
    para(
        "Book Author IAST",
        size=15,
        align="center",
        first=0,
        after=4,
        line=24.3,
        rule="exact",
    )
    para(
        "Heading 1",
        size=14,
        bold=False,
        align="center",
        first=0,
        before=48,
        after=4.3,
        line=24.3,
        rule="exact",
        keep_next=True,
    )
    para(
        "Heading 2",
        size=13,
        bold=False,
        align="left",
        first=42.5,
        before=24,
        after=4.3,
        line=24.3,
        rule="exact",
        keep_next=True,
    )
    para(
        "Center",
        size=13,
        align="center",
        first=0,
        after=4.3,
        line=20,
        rule="exact",
        keep_next=True,
    )
    para(
        "Label",
        size=12,
        align="center",
        first=0,
        after=4.3,
        line=20,
        rule="exact",
        keep_next=True,
    )
    para(
        "Label Left",
        size=12,
        align="left",
        first=42.5,
        after=4.3,
        line=20,
        rule="exact",
        keep_next=True,
    )
    para("Verse", size=12, align="center", first=0, after=24, line=24.3, rule="exact")
    para(
        "Verse IAST",
        size=11,
        align="center",
        first=0,
        after=24,
        line=24.3,
        rule="exact",
    )
    para(
        "Commentary",
        size=12,
        align="both",
        first=42.5,
        after=4.3,
        line=20,
        rule="exact",
    )
    para(
        "Signature",
        size=14,
        align="right",
        first=0,
        before=24,
        after=4.3,
        line=20,
        rule="exact",
    )
    para("Figure", align="center", first=0, after=4.3, line=1.0)
    para(
        "Contents Heading",
        size=12,
        align="center",
        first=0,
        before=0,
        after=12,
        line=20,
        rule="exact",
    )
    para(
        "Contents Item",
        size=12,
        size_cs=11,
        align="left",
        left=72,
        hanging=29.5,
        before=2,
        after=0,
        line=14.3,
        rule="exact",
        tabs=[(72, "left"), (292, "right")],
    )
    para(
        "Contents IAST",
        size=10,
        align="left",
        left=72,
        first=0,
        after=0,
        line=14.3,
        rule="exact",
    )
    folio_font, folio_size, folio_bold = "Tahoma", 8, True

# PageBreak: an almost invisible paragraph that starts a new page (markdown "* * *")
para(
    "Page Break",
    size=1,
    first=0,
    before=0,
    after=0,
    line=1,
    rule="exact",
    page_break=True,
)

# footer: centred page number (and, as in the Hindi exemplar, an empty paragraph after it)
foot = sec.footer
p = foot.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Pt(0)
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 1.0
for t in ("begin", None, "end"):
    r = p.add_run()
    r.font.name, r.font.size, r.font.bold = folio_font, Pt(folio_size), folio_bold
    set_fonts(r._r.get_or_add_rPr(), folio_font)
    if t:
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), t)
    else:
        e = OxmlElement("w:instrText")
        e.set(qn("xml:space"), "preserve")
        e.text = " PAGE "
    r._r.append(e)
if KIND == "hindi":
    q2 = foot.add_paragraph()
    q2.paragraph_format.first_line_indent = Pt(0)
    q2.paragraph_format.space_after = Pt(0)
    q2.paragraph_format.line_spacing = 1.0
    rr = q2.add_run(" ")
    rr.font.name, rr.font.size = folio_font, Pt(folio_size)

if KIND == "english":  # the English exemplar hyphenates
    s = doc.settings.element
    child(s, "w:autoHyphenation")
doc.save(str(OUT))
print("wrote", OUT)
